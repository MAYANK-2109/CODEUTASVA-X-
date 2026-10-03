import io
import openpyxl
from app.portfolio.extractor import extract_holdings_from_excel, resolve_symbol


def test_resolve_symbol():
    assert resolve_symbol("Reliance Industries Ltd") == "RELIANCE"
    assert resolve_symbol("Tata Consultancy Services") == "TCS"
    assert resolve_symbol("Infosys Limited") == "INFY"
    assert resolve_symbol("HDFC Bank") == "HDFCBANK"
    assert resolve_symbol("Tata Motors") == "TATAMOTORS"
    assert resolve_symbol("Random Stock (ABCDEF)") == "ABCDEF"


def test_excel_with_top_metadata_and_ltp():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Portfolio"
    ws.append(["Client Name: John Doe", "Client Code: 998877"])
    ws.append(["Generated on: 2026-10-03"])
    ws.append([])
    ws.append(["Symbol", "Company Name", "Qty", "Avg Buy Price", "LTP", "Current Value"])
    ws.append(["INFY", "Infosys Ltd", 15, 1420.5, 1510.0, 22650.0])
    ws.append(["TCS", "Tata Consultancy Services", 8, 3500.0, 3750.0, 30000.0])
    ws.append(["Total", "", 23, "", "", 52650.0])

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    holdings = extract_holdings_from_excel(data, "xlsx")
    assert len(holdings) == 2

    h1 = holdings[0]
    assert h1["symbol"] == "INFY"
    assert h1["name"] == "Infosys Ltd"
    assert h1["units"] == 15.0
    assert h1["buy_price"] == 1420.5
    assert h1["current_price"] == 1510.0

    h2 = holdings[1]
    assert h2["symbol"] == "TCS"
    assert h2["name"] == "Tata Consultancy Services"
    assert h2["units"] == 8.0
    assert h2["buy_price"] == 3500.0
    assert h2["current_price"] == 3750.0


def test_excel_without_symbol_column():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Particulars", "Holding Quantity", "Purchase Price"])
    ws.append(["Reliance Industries Limited", 25, 2350.0])
    ws.append(["HDFC Bank Ltd", 40, 1480.0])

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    holdings = extract_holdings_from_excel(data, "xlsx")
    assert len(holdings) == 2

    assert holdings[0]["symbol"] == "RELIANCE"
    assert holdings[0]["units"] == 25.0
    assert holdings[0]["buy_price"] == 2350.0

    assert holdings[1]["symbol"] == "HDFCBANK"
    assert holdings[1]["units"] == 40.0
    assert holdings[1]["buy_price"] == 1480.0


GROWW_STATEMENT = [
    ["Name", "Mayank Kumar Sahu"], ["Unique Client Code", "2953131503"], [],
    ["Holdings statement for stocks as on 02-10-2026"], [],
    ["Summary"], ["Invested Value", 16938.02], ["Closing Value", 13595.57], ["Unrealised P&L", -3342.45], [],
    ["Stock Name", "ISIN", "Quantity", "Average buy price", "Buy value", "Closing price", "Closing value", "Unrealised P&L"],
    ["NIP IND ETF GOLD BEES", "INF204KB17I5", 8, 135.92, 1087.36, 121.43, 971.44, -115.92],
    ["RELIANCE INDUSTRIES LTD", "INE002A01018", 1, 1311, 1311, 1167.7, 1167.7, -143.3],
    ["STATE BANK OF INDIA", "INE062A01020", 2, 963.82, 1927.64, 954.1, 1908.2, -19.44],
    ["TATAAML-TATAGOLD", "INF277KA1976", 21, 15.69, 329.49, 14.28, 299.88, -29.61],
]


def _xlsx(rows) -> bytes:
    from io import BytesIO

    import openpyxl

    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_groww_holdings_statement_skips_the_summary_block():
    from app.portfolio.extractor import extract_holdings

    rows = extract_holdings(_xlsx(GROWW_STATEMENT), "holdings.xlsx")
    assert [r["name"] for r in rows] == [
        "NIP IND ETF GOLD BEES", "RELIANCE INDUSTRIES LTD", "STATE BANK OF INDIA", "TATAAML-TATAGOLD"]
    reliance = rows[1]
    assert (reliance["units"], reliance["buy_price"], reliance["current_price"]) == (1.0, 1311.0, 1167.7)
    assert reliance["isin"] == "INE002A01018" and reliance["type"] == "STOCK"
    # Gold and commodity instruments are classified as COMMODITY
    assert rows[0]["type"] == "COMMODITY" and rows[3]["type"] == "COMMODITY"
    assert rows[3]["current_price"] == 14.28


def test_statement_text_lines_are_parsed_when_a_pdf_has_no_table():
    from app.portfolio.extractor import _extract_from_text

    text = """Holdings statement for stocks as on 02-10-2026
Invested Value 16938.02
Stock Name ISIN Quantity Average buy price Buy value Closing price Closing value Unrealised P&L
NIPPONAMC - NETFSILVER INF204KC1402 35 295.31 10335.85 208.71 7304.85 -3031
RELIANCE INDUSTRIES LTD INE002A01018 1 1,311 1,311 1,167.7 1,167.7 -143.3"""
    rows = _extract_from_text(text)
    assert [(r["name"], r["isin"], r["units"], r["buy_price"], r["current_price"]) for r in rows] == [
        ("NIPPONAMC - NETFSILVER", "INF204KC1402", 35.0, 295.31, 208.71),
        ("RELIANCE INDUSTRIES LTD", "INE002A01018", 1.0, 1311.0, 1167.7),
    ]
    assert rows[0]["type"] == "COMMODITY"


def test_upload_fills_symbol_from_isin_and_prefers_the_live_price(monkeypatch):
    from app.ingestion import prices
    from app.portfolio import router
    from app.tools import resolver

    monkeypatch.setattr(resolver, "search", lambda q: {"INE002A01018": "RELIANCE.NS", "INF204KB17I5": "GOLDBEES.NS"}.get(q))
    monkeypatch.setattr(prices, "get_latest_prices", lambda tickers: ({"RELIANCE.NS": 1180.0}, "now"))
    holdings = [
        {"name": "RELIANCE INDUSTRIES LTD", "symbol": "", "isin": "INE002A01018", "buy_price": 1311.0, "current_price": 1167.7},
        {"name": "NIP IND ETF GOLD BEES", "symbol": "", "isin": "INF204KB17I5", "buy_price": 135.92, "current_price": 121.43},
        {"name": "Unknown Co", "symbol": "", "isin": "", "buy_price": 10.0, "current_price": None},
    ]
    router._enrich(holdings)
    assert (holdings[0]["symbol"], holdings[0]["current_price"]) == ("RELIANCE", 1180.0)
    # No live quote: the symbol is still filled and the statement's closing price is kept.
    assert (holdings[1]["symbol"], holdings[1]["current_price"]) == ("GOLDBEES", 121.43)
    # Neither a live quote nor a closing price: unknown, never the buy price.
    assert (holdings[2]["symbol"], holdings[2]["current_price"]) == ("", None)


def test_wrapped_pdf_header_cells_and_names_are_normalised():
    from app.portfolio.extractor import _extract_from_tables

    table = [
        ["Stock Name", "ISIN", "Quantity", "Average buy\nprice", "Buy\nvalue", "Closing\nprice", "Closing\nvalue"],
        ["ZERODHAAMC -\nSILVERCASE", "INF0R8F01091", "81", "22.17", "1795.77", "22.12", "1791.72"],
    ]
    (row,) = _extract_from_tables([table])
    assert row["name"] == "ZERODHAAMC - SILVERCASE"
    assert (row["units"], row["buy_price"], row["current_price"]) == (81.0, 22.17, 22.12)


def test_isins_and_exact_names_resolve_from_the_exchange_list(monkeypatch):
    from app.tools import resolver

    def no_network(*args, **kwargs):
        raise AssertionError("the exchange list should answer this without a search")

    monkeypatch.setattr(resolver.yf, "Search", no_network)
    assert resolver.candidates("", "INF204KB17I5", "NIP IND ETF GOLD BEES") == ["GOLDBEES.NS"]
    assert resolver.candidates("", "INF200KA16D8", "SBI-ETF GOLD") == ["SETFGOLD.NS"]
    assert resolver.candidates("", "INE002A01018", "") == ["RELIANCE.NS"]
    assert resolver.candidates("", "", "State Bank of India") == ["SBIN.NS"]
    assert resolver.candidates("", "", "RELIANCE INDUSTRIES LTD") == ["RELIANCE.NS"]


def test_asset_class_inferences():
    from app.portfolio.extractor import _infer_type

    # Commodities (Gold, Silver, SGB, Bullion, Commodity funds)
    assert _infer_type("Nippon India ETF Gold BeES", "INF204KB17I5") == "COMMODITY"
    assert _infer_type("SBI Gold Fund - Direct Plan - Growth") == "COMMODITY"
    assert _infer_type("Nippon India Silver FoF") == "COMMODITY"
    assert _infer_type("ICICI Prudential Commodities Fund") == "COMMODITY"
    assert _infer_type("Sovereign Gold Bond 2028-29 Series I", "IN0020190394") == "COMMODITY"
    assert _infer_type("SGBMAY29") == "COMMODITY"
    assert _infer_type("Crude Oil August Future") == "COMMODITY"
    assert _infer_type("Natural Gas") == "COMMODITY"
    assert _infer_type("ZERODHAAMC - SILVERCASE") == "COMMODITY"
    assert _infer_type("TATAAML-TATAGOLD") == "COMMODITY"

    # Mutual Funds (Equity, Debt, Hybrid, Index Funds, ELSS)
    assert _infer_type("Parag Parikh Flexi Cap Fund - Direct Plan - Growth", "INF879O01019") == "MF"
    assert _infer_type("Quant Small Cap Fund - Growth") == "MF"
    assert _infer_type("HDFC Top 100 - Regular - IDCW") == "MF"
    assert _infer_type("Mirae Asset Large Cap Plan") == "MF"
    assert _infer_type("SBI Bluechip Scheme") == "MF"
    assert _infer_type("Some Mystery Scheme", "INF123456789") == "MF"

    # ETFs (Equity and Index ETFs)
    assert _infer_type("Nippon India Nifty 50 BeES ETF") == "ETF"
    assert _infer_type("Bank BeES") == "ETF"
    assert _infer_type("CPSE ETF") == "ETF"

    # Bonds & G-Secs
    assert _infer_type("7.26% GOI 2033", "IN0020230018") == "BOND"
    assert _infer_type("NABARD NCD 2028") == "BOND"
    assert _infer_type("REC Tax Free Bond") == "BOND"

    # Stocks (Equities)
    assert _infer_type("Reliance Industries Ltd", "INE002A01018") == "STOCK"
    assert _infer_type("Infosys Limited", "INE009A01021") == "STOCK"
    assert _infer_type("Tata Motors Ltd") == "STOCK"
